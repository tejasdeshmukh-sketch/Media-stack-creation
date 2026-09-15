"""Workbook/CSV writers for the pivoted output."""

from __future__ import annotations

import io
import re

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True)
MAX_WIDTH = 60
SHEET_NAME = "output_image_links"

# openpyxl refuses these outright; scraped feeds carry them often enough that a
# single stray 0x0B would otherwise blow up the whole export.
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def sanitize(frame: pd.DataFrame) -> pd.DataFrame:
    """Strip control characters Excel will not accept from every text cell."""
    cleaned = frame.copy()
    for column in cleaned.columns:
        if cleaned[column].map(lambda v: isinstance(v, str)).any():
            cleaned[column] = cleaned[column].map(
                lambda v: _ILLEGAL.sub("", v) if isinstance(v, str) else v
            )
    return cleaned


def _autofit(worksheet, frame: pd.DataFrame) -> None:
    for position, column in enumerate(frame.columns, start=1):
        widest = len(str(column))
        # astype(str) keeps NA as NA on pandas 3, so stringify per value instead.
        for value in frame[column].head(400):
            widest = max(widest, len("" if pd.isna(value) else str(value)))
        worksheet.column_dimensions[get_column_letter(position)].width = min(widest + 2, MAX_WIDTH)


def _style_header(worksheet, columns: int) -> None:
    for position in range(1, columns + 1):
        cell = worksheet.cell(row=1, column=position)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    worksheet.freeze_panes = "B2"


def to_workbook(
    frame: pd.DataFrame,
    *,
    skipped: pd.DataFrame | None = None,
    checks: pd.DataFrame | None = None,
    warnings: list[str] | None = None,
    sheet_name: str = SHEET_NAME,
) -> bytes:
    """Build the output .xlsx in memory and hand back the bytes."""
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        _sheet(writer, sanitize(frame), sheet_name)
        if skipped is not None and not skipped.empty:
            _sheet(writer, sanitize(skipped), "skipped_rows")
        if checks is not None and not checks.empty:
            _sheet(writer, sanitize(checks), "url_check")
        if warnings:
            _sheet(writer, sanitize(pd.DataFrame({"Warning": warnings})), "warnings")
    return buffer.getvalue()


def _sheet(writer, frame: pd.DataFrame, name: str) -> None:
    frame.to_excel(writer, index=False, sheet_name=name)
    worksheet = writer.sheets[name]
    _style_header(worksheet, len(frame.columns))
    _autofit(worksheet, frame)


def to_csv(frame: pd.DataFrame) -> bytes:
    return sanitize(frame).to_csv(index=False).encode("utf-8-sig")
